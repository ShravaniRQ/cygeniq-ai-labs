"""
convert_md_to_docx.py
Converts all .md files in dummy_sharepoint_docs/ into .docx Word files,
preserving folder structure (so you can directly upload to SharePoint).
"""

import os
import re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

INPUT_DIR = "dummy_sharepoint_docs"
OUTPUT_DIR = "dummy_sharepoint_docs_docx"


def add_paragraph_with_style(doc, text, style_name):
    try:
        para = doc.add_paragraph(style=style_name)
        para.add_run(text)
    except KeyError:
        # Fallback if style not found
        doc.add_paragraph(text)


def convert_md_to_docx(md_path, docx_path):
    doc = Document()

    # Set default font
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)

    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    for line in lines:
        line = line.rstrip()

        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("- ") or line.startswith("* "):
            content = line[2:]
            # Handle **bold:** pattern in bullets
            para = doc.add_paragraph(style="List Bullet")
            parts = re.split(r"(\*\*.*?\*\*)", content)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    run = para.add_run(part[2:-2])
                    run.bold = True
                else:
                    para.add_run(part)
        elif re.match(r"^\d+\. ", line):
            content = re.sub(r"^\d+\. ", "", line)
            para = doc.add_paragraph(style="List Number")
            para.add_run(content)
        elif line.startswith("`") and line.endswith("`"):
            para = doc.add_paragraph()
            run = para.add_run(line.strip("`"))
            run.font.name = "Courier New"
            run.font.size = Pt(10)
        elif line == "" or line == "---":
            doc.add_paragraph()
        else:
            # Normal paragraph — handle inline **bold**
            para = doc.add_paragraph()
            parts = re.split(r"(\*\*.*?\*\*)", line)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    run = para.add_run(part[2:-2])
                    run.bold = True
                else:
                    para.add_run(part)

    doc.save(docx_path)
    print(f"  [OK] Converted: {docx_path}")


def main():
    converted = 0
    for root, dirs, files in os.walk(INPUT_DIR):
        for filename in files:
            if filename.endswith(".md"):
                md_path = os.path.join(root, filename)
                relative = os.path.relpath(md_path, INPUT_DIR)
                docx_relative = relative.replace(".md", ".docx")
                docx_path = os.path.join(OUTPUT_DIR, docx_relative)

                os.makedirs(os.path.dirname(docx_path), exist_ok=True)
                print(f"Converting: {md_path}")
                convert_md_to_docx(md_path, docx_path)
                converted += 1

    print(f"\nDone! {converted} file(s) converted.")
    print(f"Word files saved to: {os.path.abspath(OUTPUT_DIR)}")


if __name__ == "__main__":
    main()
